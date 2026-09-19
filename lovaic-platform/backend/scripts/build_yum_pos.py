"""
Build the Yum! India POS intelligence layer.

Reads the two raw POS workbooks shipped in YumDoc/ and pre-computes a compact
JSON of real aggregates (sales, day-part demand, category/brand mix, top items,
margins, procurement, anomalies, slow-movers). The FastAPI backend then serves
these aggregates and a rule-based "Store Performance Copilot" answers plain-
language questions off them — no numbers are invented; everything traces to the
source files.

Run:  python scripts/build_yum_pos.py
Out:  app/data/yum_pos.json
"""
from __future__ import annotations

import json
import os
from collections import defaultdict
from datetime import datetime

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
REPO = os.path.dirname(os.path.dirname(BACKEND))
YUMDOC = os.path.join(REPO, "YumDoc")
OUT = os.path.join(BACKEND, "app", "data", "yum_pos.json")

RAW_FILE = os.path.join(YUMDOC, "POS level Sample Store raw data.xlsx")
SALES_FILE = os.path.join(YUMDOC, "Sample POS Sales & Inward Data.xlsx")

MONTH_ORDER = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
               "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def num(x, default=0.0):
    try:
        if x is None or x == "":
            return default
        return float(x)
    except (TypeError, ValueError):
        return default


def rupees(x: float) -> str:
    """Indian-style short rupee formatting."""
    x = float(x)
    if abs(x) >= 1e7:
        return f"₹{x/1e7:.2f} Cr"
    if abs(x) >= 1e5:
        return f"₹{x/1e5:.2f} L"
    if abs(x) >= 1e3:
        return f"₹{x/1e3:.1f}K"
    return f"₹{x:,.0f}"


def rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    for r in it:
        yield idx, r
    wb.close()


def g(idx, r, key, default=None):
    i = idx.get(key)
    if i is None or i >= len(r):
        return default
    return r[i]


# --------------------------------------------------------------------------
# 1) Primary transactional dataset — the single Mumbai store raw data (~23k lines)
# --------------------------------------------------------------------------
def build_store_raw():
    sales_by_hour = defaultdict(float)
    txn_by_hour = defaultdict(set)
    sales_by_month = defaultdict(float)
    sales_by_cat = defaultdict(float)
    qty_by_cat = defaultdict(float)
    sales_by_sub = defaultdict(float)
    sales_by_brand = defaultdict(float)
    item_sales = defaultdict(float)
    item_qty = defaultdict(float)
    item_cat = {}

    total_sales = 0.0
    total_qty = 0.0
    lines = 0
    invoices = set()
    discount_hits = 0          # selling price < mrp (offer/markdown)
    overcharge_hits = 0        # selling price > mrp (anomaly)
    discount_value = 0.0
    store_id = city = store_class = store_type = None
    date_min = date_max = None

    for idx, r in rows(RAW_FILE, "Mumbai- 1 Store Raw Data"):
        inv = g(idx, r, "Invoice_ID")
        if inv is None:
            continue
        lines += 1
        invoices.add(inv)
        store_id = store_id or g(idx, r, "Store_ID")
        city = city or g(idx, r, "Metro_City")
        store_class = store_class or g(idx, r, "Store_Class")
        store_type = store_type or g(idx, r, "Store_Type")

        qty = num(g(idx, r, "Quantity"), 1)
        amt = num(g(idx, r, "Total_Amount"))
        mrp = num(g(idx, r, "MRP"))
        sp = num(g(idx, r, "Selling_Price"))
        total_sales += amt
        total_qty += qty

        hour = g(idx, r, "Hour")
        if hour is not None and str(hour).strip() != "":
            try:
                h = int(float(hour))
                sales_by_hour[h] += amt
                txn_by_hour[h].add(inv)
            except (TypeError, ValueError):
                pass

        month = g(idx, r, "Month")
        if month:
            sales_by_month[str(month).strip()] += amt

        cat = str(g(idx, r, "Category") or "Uncategorized").strip()
        sub = str(g(idx, r, "Sub_Category") or "Other").strip()
        brand = str(g(idx, r, "Brand") or "Na").strip()
        item = str(g(idx, r, "Item_Description") or g(idx, r, "Centralized_Description") or "Item").strip()
        sales_by_cat[cat] += amt
        qty_by_cat[cat] += qty
        sales_by_sub[sub] += amt
        if brand.lower() not in ("na", "n/a", ""):
            sales_by_brand[brand] += amt
        item_sales[item] += amt
        item_qty[item] += qty
        item_cat[item] = cat

        if mrp > 0 and sp > 0:
            if sp < mrp:
                discount_hits += 1
                discount_value += (mrp - sp) * qty
            elif sp > mrp:
                overcharge_hits += 1

        d = g(idx, r, "Date")
        if isinstance(d, str) and d.strip():
            try:
                dt = datetime.strptime(d.strip(), "%d/%m/%Y")
                date_min = dt if date_min is None else min(date_min, dt)
                date_max = dt if date_max is None else max(date_max, dt)
            except ValueError:
                pass
        elif isinstance(d, datetime):
            date_min = d if date_min is None else min(date_min, d)
            date_max = d if date_max is None else max(date_max, d)

    n_inv = len(invoices)
    aov = total_sales / n_inv if n_inv else 0
    basket = total_qty / n_inv if n_inv else 0

    def top(dic, k=8):
        return sorted(dic.items(), key=lambda kv: kv[1], reverse=True)[:k]

    cat_total = sum(sales_by_cat.values()) or 1
    categories = [
        {"name": c, "value": round(v), "share": round(v / cat_total * 100, 1)}
        for c, v in top(sales_by_cat, 8)
    ]
    brands = [{"name": b, "value": round(v)} for b, v in top(sales_by_brand, 8)]
    top_items = [
        {"name": it[:38], "revenue": round(item_sales[it]), "qty": round(item_qty[it]),
         "category": item_cat.get(it, "")}
        for it in sorted(item_sales, key=lambda x: item_sales[x], reverse=True)[:12]
    ]

    # slow movers / dead-stock (waste risk): stocked (has sales rows) but very low velocity
    slow = sorted(item_qty.items(), key=lambda kv: kv[1])[:12]
    slow_movers = [
        {"name": it[:38], "qty": round(q), "revenue": round(item_sales[it]),
         "category": item_cat.get(it, "")}
        for it, q in slow
    ]

    # day-part rollup for kitchen/workforce forecasting
    hours = []
    for h in range(24):
        hours.append({
            "hour": h,
            "sales": round(sales_by_hour.get(h, 0.0)),
            "txns": len(txn_by_hour.get(h, set())),
        })
    active = [x for x in hours if x["txns"] > 0]
    peak = max(active, key=lambda x: x["sales"]) if active else {"hour": 0}

    months = [
        {"name": m, "value": round(sales_by_month[m])}
        for m in MONTH_ORDER if m in sales_by_month
    ]

    return {
        "store": {
            "id": store_id, "city": city, "class": store_class, "type": store_type,
            "date_from": date_min.strftime("%d %b %Y") if date_min else None,
            "date_to": date_max.strftime("%d %b %Y") if date_max else None,
        },
        "kpis": {
            "total_sales": round(total_sales),
            "total_sales_fmt": rupees(total_sales),
            "transactions": n_inv,
            "line_items": lines,
            "units_sold": round(total_qty),
            "aov": round(aov, 1),
            "aov_fmt": rupees(aov),
            "basket_size": round(basket, 2),
            "unique_skus": len(item_sales),
        },
        "hours": hours,
        "peak_hour": peak["hour"],
        "months": months,
        "categories": categories,
        "brands": brands,
        "top_items": top_items,
        "slow_movers": slow_movers,
        "discount": {
            "lines_discounted": discount_hits,
            "discount_value": round(discount_value),
            "discount_value_fmt": rupees(discount_value),
            "overcharge_lines": overcharge_hits,
            "discount_pct_of_lines": round(discount_hits / lines * 100, 1) if lines else 0,
        },
    }


# --------------------------------------------------------------------------
# 2) Multi-store Sales sample — city / credit / anomaly view
# --------------------------------------------------------------------------
def build_sales_sample():
    sales_by_city = defaultdict(float)
    inv_by_city = defaultdict(set)
    credit_sales = 0.0
    cash_sales = 0.0
    total = 0.0
    anomalies = 0
    n = 0
    for idx, r in rows(SALES_FILE, "Sales"):
        inv = g(idx, r, "Invoice_ID")
        if inv is None:
            continue
        n += 1
        amt = num(g(idx, r, "Total_Amount"))
        total += amt
        city = (g(idx, r, "Metro_City") or "Other").strip()
        sales_by_city[city] += amt
        inv_by_city[city].add(inv)
        cr = str(g(idx, r, "is_credit") or "").strip().upper()
        if cr == "YES":
            credit_sales += amt
        else:
            cash_sales += amt
        mrp = num(g(idx, r, "MRP"))
        sp = num(g(idx, r, "Selling_Price"))
        if mrp > 0 and sp > mrp:
            anomalies += 1

    cities = [
        {"name": c, "value": round(v), "txns": len(inv_by_city[c])}
        for c, v in sorted(sales_by_city.items(), key=lambda kv: kv[1], reverse=True)
    ]
    return {
        "lines": n,
        "total_sales": round(total),
        "cities": cities,
        "credit_share_pct": round(credit_sales / total * 100, 1) if total else 0,
        "credit_value": round(credit_sales),
        "credit_value_fmt": rupees(credit_sales),
        "price_anomalies": anomalies,
    }


# --------------------------------------------------------------------------
# 3) Inward / procurement — supplier & margin intelligence
# --------------------------------------------------------------------------
def build_inward():
    total_inward = 0.0
    by_supplier = defaultdict(float)     # keyed by po_number (proxy for supplier order)
    margin_items = []
    orders = set()
    n = 0
    for idx, r in rows(SALES_FILE, "Inward"):
        name = g(idx, r, "items.name")
        if not name:
            continue
        n += 1
        amt = num(g(idx, r, "items.total_amount"))
        total_inward += amt
        po = g(idx, r, "items.po_number")
        if po is not None:
            by_supplier[str(po)] += amt
        oid = g(idx, r, "_id")
        if oid:
            orders.add(oid)
        mrp = num(g(idx, r, "items.mrp"))
        pp = num(g(idx, r, "items.purchase_price"))
        if mrp > 0 and pp > 0:
            margin_pct = (mrp - pp) / mrp * 100
            margin_items.append({
                "name": str(name)[:34],
                "purchase": round(pp),
                "mrp": round(mrp),
                "margin_pct": round(margin_pct, 1),
                "qty": round(num(g(idx, r, "items.ordered_quantity"))),
            })

    best = sorted(margin_items, key=lambda x: x["margin_pct"], reverse=True)[:8]
    worst = sorted(margin_items, key=lambda x: x["margin_pct"])[:8]
    return {
        "lines": n,
        "orders": len(orders),
        "total_inward": round(total_inward),
        "total_inward_fmt": rupees(total_inward),
        "po_count": len(by_supplier),
        "avg_margin_pct": round(sum(m["margin_pct"] for m in margin_items) / len(margin_items), 1) if margin_items else 0,
        "best_margin": best,
        "worst_margin": worst,
    }


# --------------------------------------------------------------------------
# 4) City-level store classification (portfolio view)
# --------------------------------------------------------------------------
def build_classification():
    """This sheet has a blank leading row, so locate the header row manually."""
    wb = openpyxl.load_workbook(RAW_FILE, read_only=True, data_only=True)
    ws = wb["City Level Store Classification"]
    out = []
    for r in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in r]
        # data rows look like ['', 'Mumbai', '208', '110', '65', '33', ...]
        if len(cells) < 6:
            continue
        city = cells[1]
        if not city or city.lower() in ("cities", "store type", ""):
            continue
        try:
            out.append({
                "city": city,
                "billed_stores": int(num(cells[2])),
                "a": int(num(cells[3])),
                "b": int(num(cells[4])),
                "c": int(num(cells[5])),
            })
        except Exception:
            continue
    wb.close()
    return out


def main():
    print("Reading POS workbooks…")
    data = {
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "source_files": [os.path.basename(RAW_FILE), os.path.basename(SALES_FILE)],
        "store_raw": build_store_raw(),
        "sales_sample": build_sales_sample(),
        "inward": build_inward(),
        "classification": build_classification(),
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w") as f:
        json.dump(data, f, indent=2, default=str)
    k = data["store_raw"]["kpis"]
    print(f"Wrote {OUT}")
    print(f"  Store {data['store_raw']['store']['id']} ({data['store_raw']['store']['city']}): "
          f"{k['transactions']:,} txns · {k['total_sales_fmt']} sales · AOV {k['aov_fmt']} · "
          f"{k['unique_skus']:,} SKUs")
    print(f"  Peak hour: {data['store_raw']['peak_hour']}:00")
    print(f"  Inward: {data['inward']['total_inward_fmt']} across {data['inward']['orders']} orders, "
          f"avg margin {data['inward']['avg_margin_pct']}%")


if __name__ == "__main__":
    main()
